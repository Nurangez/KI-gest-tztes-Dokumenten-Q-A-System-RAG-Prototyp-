import pdfplumber
from docx import Document
import openpyxl
from pptx import Presentation
import csv


def load_pdf(path: str) -> str:
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def load_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def load_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            zeile = " ".join(str(cell) for cell in row if cell is not None)
            text += zeile + "\n"
    return text


def load_pptx(path: str) -> str:
    prs = Presentation(path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        text += run.text + " "
        text += "\n"
    return text


def load_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def load_csv(path: str) -> str:
    text = ""
    with open(path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            text += " ".join(row) + "\n"
    return text